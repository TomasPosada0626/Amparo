"""Hoja de revisión para el abogado: SOLO los 40 casos nuevos de v3.

POR QUE SOLO LOS NUEVOS
    Leonardo Galeano ya aprobó el corpus y los 69 casos de v2. Volver a
    ponérselos delante gasta su tiempo en lo que ya revisó y, peor, invita a
    cambiar un caso viejo: v3 es SUPERCONJUNTO ESTRICTO de v2 y el puente de
    comparabilidad con las doce corridas anteriores depende de que los 69 no
    se toquen ni una coma.

QUE SE LE PIDE QUE REVISE, Y QUE NO
    El MECANISMO no se le pide: lo deriva `reglas.decidir()` del árbol del
    art. 6 del Decreto 2591 a partir de los hechos codificados, y
    `data/ampliar_eval_set.py` rechaza cualquier caso donde el motor no
    coincida con lo esperado. Lo que necesita juicio profesional es otra cosa:

      1. `hechos_ok`     ¿la codificación de los hechos corresponde al relato?
                         Es la pregunta decisiva: si los hechos están mal
                         codificados, el motor deriva una etiqueta plausible
                         y equivocada, y nada más en el proyecto lo detecta.
      2. `derecho_ok`    ¿el derecho nombrado es el que está en juego?
      3. `normas_ok`     ¿las normas citadas vienen al caso?
      4. `realista`      ¿el relato se parece a algo que llegaría de verdad?
      5. `observacion`   lo que haya que corregir, en sus palabras.

    Las cuatro primeras son listas cerradas (si / no / dudoso) para que un
    error de teclado no se cuente como desacuerdo.

    Uso: python -I data/revision_abogado_v3.py           # escribe el formulario
         python -I data/revision_abogado_v3.py --leer    # lo lee de vuelta

LEER DE VUELTA NO ES UN EXTRA
    El mismo script que escribe el formulario lo lee, como hace
    `data/doble_anotacion.py`. Si la lectura viviera en otro sitio, las
    columnas podrian dejar de corresponder sin que nada fallara. Y el
    resultado se publica como artefacto JSON: una revision profesional que se
    queda en un .xlsx en el escritorio no es parte de la entrega, y esta es la
    PRIMERA validacion humana externa del gold que tiene el proyecto --- la
    del juez sigue abierta, que es otra cosa.
"""
import json
import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "data"))
sys.path.insert(0, str(RAIZ / "src"))
import formularios  # noqa: E402

V2 = RAIZ / "data" / "eval_set_v2.json"
V3 = RAIZ / "data" / "eval_set_v3.json"
SALIDA = RAIZ / "data" / "revision_abogado_v3"

COLUMNAS_PERSONA = ("hechos_ok", "derecho_ok", "normas_ok", "realista",
                    "observacion")
CABECERA = ["id", "tipo", "derecho", "mecanismo_derivado", "relato",
            "hechos_codificados", "normas_citadas", "razon_propuesta",
            *COLUMNAS_PERSONA]
SI_NO = ["si", "no", "dudoso"]

INSTRUCCIONES = [
    "REVISION DE LOS 40 CASOS NUEVOS DEL EVAL SET v3",
    "",
    "Responsable del corpus: Leonardo Galeano, abogado.",
    "",
    "Los 69 casos que ya aprobo NO estan aqui a proposito: v3 los contiene",
    "sin un solo cambio, y la comparacion con las corridas anteriores",
    "depende de que sigan intactos.",
    "",
    "EL MECANISMO NO SE REVISA. Lo deriva el motor de reglas del arbol del",
    "art. 6 del Decreto 2591 de 1991 a partir de la columna",
    "`hechos_codificados`, y un script rechaza el caso si el motor no emite",
    "lo esperado. La columna `mecanismo_derivado` esta solo para que vea que",
    "salio.",
    "",
    "LO QUE SI NECESITA SU JUICIO, por columna:",
    "",
    "  hechos_ok    ¿La codificacion de `hechos_codificados` corresponde a lo",
    "               que cuenta el relato? ES LA PREGUNTA DECISIVA: si los",
    "               hechos estan mal codificados, el motor deriva una",
    "               etiqueta plausible y equivocada, y nada mas en el",
    "               proyecto lo detecta.",
    "  derecho_ok   ¿El derecho de la columna `derecho` es el que esta en",
    "               juego en el relato?",
    "  normas_ok    ¿Las normas de `normas_citadas` vienen al caso?",
    "  realista     ¿El relato se parece a algo que llegaria de verdad a una",
    "               consulta? Son relatos sinteticos y el proyecto lo declara;",
    "               lo que importa es que no sean inverosimiles.",
    "  observacion  Lo que haya que corregir, en sus palabras. Si marca `no`",
    "               o `dudoso` en algo, aqui es donde dice por que.",
    "",
    "Las cuatro primeras columnas son listas cerradas: si / no / dudoso.",
    "No escriba nada en las columnas de la izquierda.",
]


def leer() -> int:
    """Lee el formulario devuelto y publica el resultado como artefacto."""
    from openpyxl import load_workbook

    ruta = SALIDA.with_suffix(".xlsx")
    if not ruta.exists():
        print(f"no existe {ruta.name}: no hay nada que leer")
        return 1
    ws = load_workbook(ruta)["formulario"]
    fila_cab = next(r for r in range(1, ws.max_row + 1)
                    if ws.cell(r, 1).value == "id")
    cols = {ws.cell(fila_cab, c).value: c
            for c in range(1, ws.max_column + 1)}

    import collections
    por_caso, cuenta, observaciones = {}, {}, []
    for k in COLUMNAS_PERSONA[:4]:
        cuenta[k] = collections.Counter()
    for r in range(fila_cab + 1, ws.max_row + 1):
        cid = ws.cell(r, cols["id"]).value
        if not cid:
            continue
        d = {}
        for k in COLUMNAS_PERSONA:
            v = str(ws.cell(r, cols[k]).value or "").strip()
            d[k] = v.lower() if k != "observacion" else v
            if k in cuenta:
                cuenta[k][d[k] or "(sin responder)"] += 1
        por_caso[cid] = d
        if d["observacion"]:
            observaciones.append({"id": cid, "texto": d["observacion"]})

    # `sin_responder` NO es un `si`. Si una celda quedo vacia, el caso no esta
    # revisado y decirlo es mejor que contarlo como aprobado.
    reparos = sorted(cid for cid, d in por_caso.items()
                     if any(d[k] not in ("si",) for k in COLUMNAS_PERSONA[:4]))
    out = {
        "_artefacto": "revision_abogado_v3.json",
        "revisor": "Leonardo Galeano, abogado",
        "papel": "responsable del corpus; aprobo el corpus y los 69 casos de v2",
        "que_se_reviso": f"los {len(por_caso)} casos nuevos de eval_set_v3",
        "que_NO_se_reviso": (
            "el mecanismo, que lo deriva reglas.decidir() del arbol del art. 6 "
            "del Decreto 2591 y que un script rechaza si no coincide con lo "
            "esperado; y los 69 casos de v2, ya aprobados y no modificados"),
        "columnas": list(COLUMNAS_PERSONA),
        "conteo": {k: dict(v) for k, v in cuenta.items()},
        "casos_con_reparo": reparos,
        "n_casos_con_reparo": len(reparos),
        "observaciones": observaciones,
        "por_caso": por_caso,
        "lectura": (
            "`hechos_ok` es la columna decisiva: el motor de reglas confia en "
            "la codificacion de los hechos, asi que un relato que no la "
            "sostiene deja el gold sin respaldo y ninguna otra comprobacion "
            "del proyecto lo detecta. Un `dudoso` no dice que la etiqueta este "
            "mal: dice que el relato no permite confirmarla."),
        "alcance": (
            "Esto valida el GOLD, no el juez. La validacion del juez contra "
            "criterio humano independiente (M2-1b) sigue abierta y es otra "
            "cosa: aqui se revisaron las etiquetas, no las notas."),
    }
    (RAIZ / "data" / "revision_abogado_v3.json").write_text(
        json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")

    print("=" * 76)
    print("  REVISION DEL ABOGADO · resultado")
    print("=" * 76)
    print(f"  casos revisados: {len(por_caso)}")
    for k, v in cuenta.items():
        print(f"    {k:12s} " + "  ".join(f"{a}={b}" for a, b in v.most_common()))
    print()
    if reparos:
        print(f"  {len(reparos)} caso(s) CON REPARO: {', '.join(reparos)}")
        for o in observaciones:
            print(f"    {o['id']}: {o['texto'][:150]}")
    else:
        print("  sin reparos: las cuatro columnas en `si` en todos los casos")
    print()
    print("  escrito: data/revision_abogado_v3.json")
    print("=" * 76)
    return 0


def main() -> int:
    if "--leer" in sys.argv:
        return leer()
    v2_ids = {c["id"] for c in json.loads(V2.read_text("utf-8"))["ejemplos"]}
    v3 = json.loads(V3.read_text("utf-8"))["ejemplos"]
    nuevos = [c for c in v3 if c["id"] not in v2_ids]

    filas = []
    for c in nuevos:
        hechos = ", ".join(f"{k}={v}" for k, v in sorted(c["hechos"].items()))
        normas = "; ".join(c.get("fundamento_del_derecho") or []) or "(ninguna)"
        filas.append([c["id"], c["tipo"], c["derecho"], c["mecanismo"],
                      c["relato"], hechos, normas, c["razon"],
                      *[""] * len(COLUMNAS_PERSONA)])

    print("=" * 76)
    print("  HOJA DE REVISION PARA EL ABOGADO · casos nuevos de v3")
    print("=" * 76)
    print(f"  v2: {len(v2_ids)} casos ya aprobados, NO incluidos")
    print(f"  v3: {len(v3)} casos · {len(nuevos)} nuevos a revisar")
    print()
    import collections
    for campo in ("derecho", "tipo"):
        cuenta = collections.Counter(c[campo] for c in nuevos)
        print(f"  nuevos por {campo}: "
              + ", ".join(f"{k} {v}" for k, v in cuenta.most_common()))
    print()

    hizo_xlsx = formularios.escribir_xlsx(
        SALIDA.with_suffix(".xlsx"), INSTRUCCIONES, CABECERA, filas,
        COLUMNAS_PERSONA,
        opciones={c: SI_NO for c in COLUMNAS_PERSONA[:4]})
    if hizo_xlsx:
        print(f"  escrito: data/{SALIDA.name}.xlsx  (con desplegables)")
    else:
        print("  sin openpyxl: no se pudo escribir el .xlsx")

    # El CSV va SIEMPRE, no como respaldo del xlsx sino porque es el que se
    # puede leer sin Excel. Se escribe con BOM: sin el, Excel abre un CSV
    # UTF-8 con los acentos partidos, y ya paso una vez que el BOM acabo
    # siendo la cabecera y todas las respuestas se leyeron como `None`.
    import csv
    with SALIDA.with_suffix(".csv").open("w", encoding="utf-8-sig",
                                         newline="") as f:
        w = csv.writer(f)
        w.writerow(CABECERA)
        w.writerows(filas)
    print(f"  escrito: data/{SALIDA.name}.csv   (utf-8 con BOM, para Excel)")
    print()
    print("  Que haga el abogado: marcar las cinco columnas de la derecha y")
    print("  devolver el archivo. Un `no` en `hechos_ok` es el unico hallazgo")
    print("  que obliga a rehacer el caso antes de correr.")
    print("=" * 76)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
